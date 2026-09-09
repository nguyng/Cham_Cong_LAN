"""
Tạo lại Module_ChamCong_LAN.docx đồng bộ với Module_ChamCong_LAN.md
"""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

doc = Document()

# ── Font mặc định ──────────────────────────────────────────────────────────────
style = doc.styles['Normal']
style.font.name = 'Times New Roman'
style.font.size = Pt(12)

# ── Tiện ích ───────────────────────────────────────────────────────────────────
def heading(text, level=1, center=False):
    h = doc.add_heading(text, level=level)
    if center:
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for r in h.runs:
        r.font.name = 'Times New Roman'
    return h

def para(text='', bold=False, italic=False, indent=0):
    p = doc.add_paragraph()
    if indent:
        p.paragraph_format.left_indent = Inches(indent * 0.3)
    if text:
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(12)
        r.bold = bold
        r.italic = italic
    return p

def bullet(text, level=1):
    p = doc.add_paragraph(style='List Bullet')
    p.paragraph_format.left_indent = Inches(level * 0.3)
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(12)
    return p

def info_line(label, value):
    p = doc.add_paragraph()
    rb = p.add_run(label)
    rb.bold = True
    rb.font.name = 'Times New Roman'
    rv = p.add_run(value)
    rv.font.name = 'Times New Roman'
    rv.font.size = Pt(12)

def section_note(text):
    """Dòng chú thích nhỏ (in nghiêng) dưới heading module."""
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.italic = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(11)
    return p

def add_old_new(old_text, new_bullets):
    """Thêm cặp Quy trình cũ / Quy trình mới."""
    p_old = doc.add_paragraph()
    r_old = p_old.add_run('Quy trình cũ (Thủ công): ')
    r_old.bold = True
    r_old.font.name = 'Times New Roman'
    r_old.font.size = Pt(12)
    r_body = p_old.add_run(old_text)
    r_body.font.name = 'Times New Roman'
    r_body.font.size = Pt(12)

    p_new = doc.add_paragraph()
    r_new = p_new.add_run('Quy trình mới (Tin học hóa):')
    r_new.bold = True
    r_new.font.name = 'Times New Roman'
    r_new.font.size = Pt(12)

    for b in new_bullets:
        bullet(b)

def separator():
    doc.add_paragraph()

# ══════════════════════════════════════════════════════════════════════════════
# TIÊU ĐỀ CHÍNH
# ══════════════════════════════════════════════════════════════════════════════
h0 = doc.add_heading('PHẦN MỀM CHẤM CÔNG TRÊN MẠNG LAN', level=0)
h0.alignment = WD_ALIGN_PARAGRAPH.CENTER
for r in h0.runs:
    r.font.name = 'Times New Roman'
    r.font.size = Pt(16)

heading('Mô Tả Các Module Chức Năng', level=1, center=True)
separator()

info_line('Môn học: ', 'Lập Trình Mạng')
info_line('Kiến trúc: ', 'Microservices — Spring Boot (Java) + Python FastAPI + React')
info_line('Môi trường: ', 'Mạng nội bộ LAN (không cần Internet)')
separator()

# Bảng ký hiệu
p_ky = doc.add_paragraph()
r_ky = p_ky.add_run('Bảng ký hiệu công nghệ:')
r_ky.bold = True; r_ky.font.name = 'Times New Roman'

tbl_ky = doc.add_table(rows=4, cols=2)
tbl_ky.style = 'Table Grid'
headers_ky = [['Ký hiệu', 'Ý nghĩa'],
              ['☕ [Java]',   'Tính năng do Spring Boot thực thi — Member 1 hoặc 3 phụ trách'],
              ['🐍 [Python]', 'Tính năng do Python FastAPI thực thi — Member 2 phụ trách'],
              ['⚛️ [React]',  'Tính năng do React thực thi — Member 4 hoặc 5 phụ trách']]
for i, row_data in enumerate(headers_ky):
    row = tbl_ky.rows[i].cells
    for j, val in enumerate(row_data):
        row[j].text = val
        for pr in row[j].paragraphs:
            for rr in pr.runs:
                rr.font.name = 'Times New Roman'
                rr.font.size = Pt(11)
                if i == 0:
                    rr.bold = True
separator()

# ══════════════════════════════════════════════════════════════════════════════
# MODULE 1
# ══════════════════════════════════════════════════════════════════════════════
heading('Module 1: Chấm Công (Attendance)', level=1)
section_note('🐍 Thực thi bởi: Python FastAPI (Member 2)')
para('Trung tâm của hệ thống — ghi nhận toàn bộ hoạt động vào/ra của nhân viên theo thời gian thực.')
separator()

heading('1.1. Chấm Công Vào / Ra 🐍', level=2)
para('1.1.1. Nhân viên đăng nhập và bấm nút "Chấm Công Vào" — Python API ghi nhận giờ vào, địa chỉ IP máy tính, lưu vào bảng attendance_logs.')
para('1.1.2. Cuối ca, nhân viên bấm "Chấm Công Ra" — Python API tính tổng số giờ làm trong ngày và cập nhật bản ghi.')
para('1.1.3. Tự động ghi nhận trạng thái: Đúng giờ / Đi muộn / Về sớm bằng cách so sánh giờ chấm công với ca làm việc lấy từ MySQL.')

heading('1.2. Giới Hạn IP Chấm Công 🐍 (Điểm khác biệt)', level=2)
para('1.2.1. Python kiểm tra IP của request: chỉ chấp nhận IP nằm trong dải 192.168.x.x (nội bộ công ty).')
para('1.2.2. Nếu IP bên ngoài → API trả về lỗi 403, ghi log cảnh báo vào hệ thống.')
para('1.2.3. Admin có thể cấu hình whitelist IP qua trang Admin (gọi API Spring Boot để lưu cấu hình).')

heading('1.3. Random Check — Xác Nhận Hiện Diện Ngẫu Nhiên 🐍 (Điểm khác biệt)', level=2)
para('1.3.1. Python dùng APScheduler lập lịch gửi thông báo bất ngờ đến màn hình nhân viên trong khung giờ làm việc (qua WebSocket).')
para('1.3.2. React hiển thị popup đếm ngược — nhân viên phải bấm "Xác nhận tôi đang làm việc" trong vòng X phút.')
para('1.3.3. Nếu không phản hồi → Python tự động cảnh báo Manager qua WebSocket, ghi nhận vào lịch sử.')
para('1.3.4. Lịch sử các lần Random Check (thời gian, có phản hồi hay không) được lưu lại để Manager xem xét.')

heading('1.4. Phát Hiện & Gắn Cờ Bất Thường 🐍 (Điểm khác biệt)', level=2)
para('1.4.1. Python tự động gắn cờ ⚠️ các ca chấm công đáng ngờ: vào/ra cách nhau dưới 10 phút, chấm công ngoài giờ quy định.')
para('1.4.2. Danh sách ca bất thường hiển thị trên Dashboard Manager để kiểm tra cuối ngày/tuần.')
separator()

# ══════════════════════════════════════════════════════════════════════════════
# MODULE 2
# ══════════════════════════════════════════════════════════════════════════════
heading('Module 2: Quản Lý Nhân Sự (Human Resource Management)', level=1)
section_note('☕ Thực thi bởi: Spring Boot (Member 1)')
para('Tổ chức cơ cấu nhân sự, phân công công việc và quản lý thông tin nhân viên.')
separator()

heading('2.1. Đăng Nhập & Phân Quyền ☕', level=2)
para('2.1.1. Spring Boot xử lý đăng nhập, tạo JWT Token có chứa thông tin vai trò (role).')
para('2.1.2. 3 cấp quyền: Employee → Manager → Admin. Mỗi cấp chỉ thấy và dùng được chức năng của mình.')
para('2.1.3. JWT Token được React lưu vào localStorage và gửi kèm trong mọi request đến cả Spring Boot lẫn Python FastAPI.')

heading('2.2. Quản Lý Nhân Viên ☕', level=2)
para('2.2.1. Thêm / Sửa / Xóa thông tin nhân viên (Họ tên, Mã NV, Email, Số điện thoại, Phòng ban, Vai trò).')
para('2.2.2. Tìm kiếm và lọc nhân viên theo phòng ban, ca làm việc, trạng thái (đang làm / đã nghỉ).')
para('2.2.3. Mỗi nhân viên mới được tự động tạo tài khoản đăng nhập với mật khẩu mặc định.')

heading('2.3. Quản Lý Phòng Ban ☕', level=2)
para('2.3.1. Tạo / Sửa / Xóa các phòng ban trong công ty.')
para('2.3.2. Gán nhân viên vào phòng ban, gán Manager phụ trách từng phòng.')
para('2.3.3. Xem danh sách nhân viên theo từng phòng ban, kèm trạng thái có mặt hôm nay.')

heading('2.4. Quản Lý Ca Làm Việc ☕', level=2)
para('2.4.1. Tạo các ca làm việc: Ca sáng, Ca chiều, Ca tối, Ca hành chính... với giờ bắt đầu/kết thúc cụ thể.')
para('2.4.2. Phân ca cho từng nhân viên theo tuần hoặc theo tháng.')
para('2.4.3. Xem lịch phân ca dưới dạng bảng (theo tuần/tháng). Python sẽ đọc bảng này để tính đúng giờ / muộn.')
separator()

# ══════════════════════════════════════════════════════════════════════════════
# MODULE 3
# ══════════════════════════════════════════════════════════════════════════════
heading('Module 3: Quản Lý Nghỉ Phép (Leave Management)', level=1)
section_note('☕ Thực thi bởi: Spring Boot (Member 3) — Tự động hóa luồng: Xin phép → Duyệt → Cập nhật chuyên cần.')
separator()

heading('3.1. Xin Nghỉ Phép ☕ + ⚛️', level=2)
para('3.1.1. Nhân viên tạo đơn xin nghỉ: chọn loại nghỉ (Nghỉ phép, Nghỉ ốm, Nghỉ không lương...), ngày bắt đầu, ngày kết thúc, lý do.')
para('3.1.2. Spring Boot lưu đơn, gửi thông báo đến Manager phụ trách qua WebSocket của Python (real-time).')
para('3.1.3. React hiển thị trạng thái đơn: Chờ duyệt / Đã duyệt / Bị từ chối.')

heading('3.2. Duyệt Nghỉ Phép ☕ + ⚛️', level=2)
para('3.2.1. Manager xem danh sách đơn nghỉ phép của phòng đang chờ duyệt trên React.')
para('3.2.2. Chấp thuận hoặc từ chối kèm lý do — Spring Boot xử lý, nhân viên nhận thông báo ngay lập tức qua WebSocket.')
para('3.2.3. Khi duyệt, Spring Boot tự động trừ số ngày phép còn lại của nhân viên đó.')

heading('3.3. Theo Dõi Ngày Phép ☕', level=2)
para('3.3.1. Mỗi nhân viên xem số ngày phép còn lại trong năm qua giao diện React.')
para('3.3.2. Admin cấu hình số ngày phép tối đa theo chính sách công ty.')
para('3.3.3. Cảnh báo khi nhân viên đã dùng hết hoặc gần hết ngày phép.')
separator()

# ══════════════════════════════════════════════════════════════════════════════
# MODULE 4
# ══════════════════════════════════════════════════════════════════════════════
heading('Module 4: Báo Cáo & Thống Kê (Reports & Analytics)', level=1)
section_note('🐍 + ☕ Thực thi bởi: Python (xuất file) + Spring Boot (API tổng hợp) — Member 2 & 3 phụ trách.')
para('Tổng hợp dữ liệu chuyên cần — đầu vào để tính lương và đánh giá hiệu suất.')
separator()

heading('4.1. Dashboard Real-time ⚛️ + 🐍', level=2)
para('4.1.1. React kết nối WebSocket của Python — hiển thị ai đang có mặt tại thời điểm hiện tại, cập nhật ngay khi có người chấm công.')
para('4.1.2. Dashboard cảnh báo: danh sách nhân viên không phản hồi Random Check trong ngày.')
para('4.1.3. Bộ lọc nhanh: xem theo phòng ban, theo ca làm việc.')

heading('4.2. Báo Cáo Theo Ngày ☕', level=2)
para('4.2.1. Spring Boot tổng hợp: danh sách có mặt / vắng mặt / đi muộn trong ngày từ bảng attendance_logs.')
para('4.2.2. Xem chi tiết giờ vào/ra của từng người, trạng thái Random Check.')
para('4.2.3. Lọc theo phòng ban hoặc cá nhân.')

heading('4.3. Báo Cáo Theo Tháng & Xuất File 🐍', level=2)
para('4.3.1. Python tổng hợp: số ngày công, số giờ làm, số lần đi muộn, số lần không phản hồi Random Check.')
para('4.3.2. Xuất file Excel (.xlsx) dùng thư viện openpyxl — để tính lương cuối tháng.')
para('4.3.3. Xuất file PDF dùng thư viện reportlab — bảng tổng hợp ngày công chính thức.')
separator()

# ══════════════════════════════════════════════════════════════════════════════
# MODULE 5
# ══════════════════════════════════════════════════════════════════════════════
heading('Module 5: Hệ Thống & Bảo Mật (System & Security)', level=1)
section_note('☕ + 🐍 Thực thi bởi: Cả hai backend — Nền tảng vận hành toàn bộ hệ thống trên mạng LAN.')
separator()

heading('5.1. Bảo Mật JWT Dùng Chung ☕ + 🐍 (Quan trọng)', level=2)
para('5.1.1. Spring Boot tạo JWT Token khi đăng nhập, ký bằng secret key chung.')
para('5.1.2. Python FastAPI xác minh JWT Token từ React bằng cùng secret key đó trước khi xử lý mọi request.')
para('5.1.3. Token hết hạn sau thời gian quy định, React tự động redirect về trang đăng nhập.')

heading('5.2. Triển Khai Microservices Trên Mạng LAN 🐍 + ☕ (Điểm khác biệt)', level=2)
para('5.2.1. 1 máy chủ duy nhất chạy đồng thời: Spring Boot (cổng 8080) + Python FastAPI (cổng 8000) + MySQL.')
para('5.2.2. Toàn bộ nhân viên truy cập React qua IP nội bộ — không cần cài đặt gì trên máy client.')
para('5.2.3. Không cần Internet, không phụ thuộc dịch vụ bên ngoài — dữ liệu hoàn toàn nội bộ.')
para('5.2.4. File start_server.bat tự động khởi động toàn bộ hệ thống chỉ bằng 1 cú double-click.')

heading('5.3. WebSocket Real-time 🐍', level=2)
para('5.3.1. Python FastAPI quản lý toàn bộ kênh WebSocket — broadcast khi có chấm công mới, khi Random Check kích hoạt, khi có đơn nghỉ phép.')
para('5.3.2. Dashboard Manager cập nhật tức thì, không cần F5 trang.')
para('5.3.3. Popup Random Check hiện trên màn hình nhân viên ngay lập tức khi Python gửi lệnh.')

heading('5.4. Nhật Ký Hệ Thống (Audit Log) ☕', level=2)
para('5.4.1. Spring Boot ghi lại toàn bộ hoạt động quan trọng: đăng nhập, thay đổi nhân sự, duyệt đơn.')
para('5.4.2. Python ghi riêng log chấm công và Random Check.')
para('5.4.3. Admin xem nhật ký hợp nhất để kiểm tra bảo mật và truy vết sự cố.')
separator()

# ══════════════════════════════════════════════════════════════════════════════
# BẢNG TỔNG HỢP MODULE
# ══════════════════════════════════════════════════════════════════════════════
heading('Tóm Tắt — Ma Trận Module', level=1)

tbl = doc.add_table(rows=6, cols=4)
tbl.style = 'Table Grid'
headers = ['Module', 'Công nghệ', 'Thành viên', 'Điểm độc đáo']
for i, h_text in enumerate(headers):
    tbl.rows[0].cells[i].text = h_text
    for pr in tbl.rows[0].cells[i].paragraphs:
        for rr in pr.runs:
            rr.bold = True; rr.font.name = 'Times New Roman'

rows_data = [
    ('Module 1 — Chấm Công',   '🐍 Python FastAPI',       'Member 2',     '🌟 Giới hạn IP, Random Check, Gắn cờ bất thường'),
    ('Module 2 — Nhân Sự',     '☕ Spring Boot (Java)',    'Member 1',     '➖ Học từ HRM thương mại, nền tảng bảo mật tốt'),
    ('Module 3 — Nghỉ Phép',   '☕ Spring Boot (Java)',    'Member 3',     '🌟 Thông báo real-time qua WebSocket Python'),
    ('Module 4 — Báo Cáo',     '🐍 Python + ☕ Java',      'Member 2 & 3', '🌟 Báo cáo Random Check, xuất Excel/PDF bằng Python'),
    ('Module 5 — Hệ Thống',    '🐍 + ☕ Cả hai',           'Member 1,2,3', '🌟 Microservices LAN, JWT chung, start_server.bat'),
]
for i, (m, tech, member, note) in enumerate(rows_data):
    row = tbl.rows[i+1].cells
    for j, val in enumerate([m, tech, member, note]):
        row[j].text = val
        for pr in row[j].paragraphs:
            for rr in pr.runs:
                rr.font.name = 'Times New Roman'
                rr.font.size = Pt(11)
separator()

# ══════════════════════════════════════════════════════════════════════════════
# PHẦN 9: QUY TRÌNH TIN HỌC HÓA (AS-IS vs TO-BE)
# ══════════════════════════════════════════════════════════════════════════════
heading('9. Quy Trình Tin Học Hóa Nghiệp Vụ Chấm Công (As-Is vs To-Be)', level=1)
para('Quy trình này mô tả cách hệ thống phần mềm thay thế và tối ưu hóa các quy trình quản lý nhân sự — chấm công thủ công truyền thống tại doanh nghiệp.')
separator()

# 9.1
heading('9.1. Quản lý Hồ sơ & Phân ca làm việc', level=2)
add_old_new(
    'Bộ phận HR dùng sổ sách hoặc file Excel rời rạc để lưu thông tin nhân viên. Trưởng phòng phải xếp lịch ca làm thủ công trên bảng trắng hoặc gửi qua Zalo/Email, dễ xảy ra nhầm lẫn, trùng ca.',
    [
        'Admin khai báo toàn bộ nhân viên và cơ cấu phòng ban trên Hệ thống (Spring Boot).',
        'Manager thiết lập ca làm việc và gán ca cho nhân viên bằng vài click chuột trên giao diện Web.',
        'Dữ liệu đồng bộ tức thời đến mọi nhân viên.',
    ]
)
separator()

# 9.2
heading('9.2. Hoạt động Chấm Công Hàng Ngày', level=2)
add_old_new(
    'Nhân viên phải xếp hàng quét vân tay tại máy gắn ở cửa ra vào. Thường xuyên gặp lỗi (máy không nhận vân tay, tay ướt, quên quét) và phải nhờ HR xử lý ngoại lệ.',
    [
        'Nhân viên ngồi ngay tại bàn làm việc, mở trình duyệt và click "Chấm Công Vào/Ra".',
        'Trạm Python FastAPI tự động đánh mốc thời gian thực, đối chiếu IP mạng LAN (không cho phép dùng 4G/Wifi ngoài) để xác thực vị trí.',
        'Tự động phân loại: Đúng giờ / Đi muộn / Về sớm.',
    ]
)
separator()

# 9.3
heading('9.3. Giám sát & Chống gian lận (Trong giờ hành chính)', level=2)
add_old_new(
    'Quản lý phải đi vòng quanh văn phòng giám sát, hoặc kiểm tra camera để xem nhân viên có mặt tại bàn hay không (tốn thời gian, không hiệu quả). Máy vân tay không thể chống được việc "nhờ người quẹt thẻ hộ rồi đi ra ngoài".',
    [
        'Hệ thống áp dụng AI/Engine Random Check (Python APScheduler), thỉnh thoảng sẽ tự động bật Popup trên màn hình máy tính của nhân viên.',
        'Nhân viên phải click xác nhận trong vòng X phút.',
        'Nếu bỏ lỡ → hệ thống tự động ghi nhận vắng mặt và bắn thông báo Real-time qua WebSocket về Dashboard của Manager.',
    ]
)
separator()

# 9.4
heading('9.4. Xử lý Đơn xin nghỉ phép', level=2)
add_old_new(
    'Nhân viên viết đơn giấy hoặc gửi email. Đơn phải qua nhiều tay (Trưởng nhóm → Trưởng phòng → HR), mất thời gian chờ đợi, dễ thất lạc đơn và khó theo dõi quỹ ngày phép còn lại.',
    [
        'Nhân viên điền Form xin nghỉ ngay trên Web (React).',
        'Manager nhận được thông báo, duyệt/từ chối chỉ bằng 1 click.',
        'Hệ thống (Java) tự động cập nhật trạng thái và trừ dần vào quỹ ngày phép của nhân viên đó.',
    ]
)
separator()

# 9.5
heading('9.5. Tổng hợp Báo Cáo & Tính Lương Cuối Tháng', level=2)
add_old_new(
    'Kế toán/HR mất từ 2–3 ngày cuối tháng để trút dữ liệu từ máy vân tay ra USB, ngồi mò mẫm đối chiếu file Excel vân tay với đống giấy tờ xin nghỉ phép để chốt công. Cực kỳ dễ sai sót.',
    [
        'Hệ thống tự động tích hợp dữ liệu chấm công (Python) và dữ liệu nghỉ phép hợp lệ (Java).',
        'HR chỉ việc ấn nút "Xuất báo cáo tháng", hệ thống sẽ xuất ra file Excel/PDF (bằng thư viện openpyxl / reportlab) hoàn chỉnh chỉ trong 3–5 giây.',
        'Dữ liệu chính xác tuyệt đối.',
    ]
)
separator()

# ══════════════════════════════════════════════════════════════════════════════
# BẢNG SO SÁNH TÓM TẮT (Slide thuyết trình)
# ══════════════════════════════════════════════════════════════════════════════
heading('Bảng So Sánh Tóm Tắt — Dùng cho Slide thuyết trình', level=1)
para('Tóm tắt phần quy trình tin học hóa thành bảng so sánh ngắn gọn:', italic=True)
separator()

tbl2 = doc.add_table(rows=5, cols=4)
tbl2.style = 'Table Grid'
headers2 = ['Nghiệp Vụ', 'Trước đây (Thủ công)', 'Hệ thống LAN System (Mới)', 'Lợi ích mang lại']
for i, h_text in enumerate(headers2):
    tbl2.rows[0].cells[i].text = h_text
    for pr in tbl2.rows[0].cells[i].paragraphs:
        for rr in pr.runs:
            rr.bold = True; rr.font.name = 'Times New Roman'; rr.font.size = Pt(11)

rows2 = [
    ('Điểm danh', 'Máy vân tay / Sổ sách',            'Click trên máy tính (Check IP LAN)',       'Nhanh chóng, không ùn tắc, 0đ chi phí máy móc'),
    ('Giám sát',  'Bằng mắt thường / Camera',          'Popup Random Check + Gắn cờ tự động',     'Chủ động chống gian lận, cảnh báo real-time'),
    ('Nghỉ phép', 'Đơn giấy, ký nháy',                 'Form điện tử, Duyệt 1-click',              'Tiết kiệm thời gian, không thất lạc giấy tờ'),
    ('Chốt công', 'Mất 2–3 ngày đối chiếu Excel',      'Xuất báo cáo tự động trong 5 giây',        'Chính xác 100%, giảm tải cho HR'),
]
for i, row_data in enumerate(rows2):
    row = tbl2.rows[i+1].cells
    for j, val in enumerate(row_data):
        row[j].text = val
        for pr in row[j].paragraphs:
            for rr in pr.runs:
                rr.font.name = 'Times New Roman'
                rr.font.size = Pt(11)

# ── Lưu ────────────────────────────────────────────────────────────────────────
doc.save(r'd:\LapTrinhMang\Module_ChamCong_LAN.docx')
print('Da luu thanh cong: Module_ChamCong_LAN.docx')
